import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Documentación de Integraciones y Arreglos — Valhalla SOC', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata
    p = doc.add_paragraph()
    p.add_run('Autor: ').bold = True
    p.add_run('Rosalino Martinez\n')
    p.add_run('Fecha: ').bold = True
    p.add_run('8 de mayo de 2026')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Section 1
    doc.add_heading('1. Integración de Inteligencia Artificial (Ollama)', level=1)
    doc.add_paragraph('Se ha implementado un sistema de análisis de amenazas basado en IA local para reducir la carga de los analistas y proporcionar contextos enriquecidos.', style='List Bullet')
    doc.add_paragraph('Script de Integración (custom-ollama.py): Desarrollo de un componente personalizado para Wazuh que envía alertas críticas a modelos de lenguaje para su interpretación.', style='List Bullet')
    doc.add_paragraph('Análisis Forense Automatizado: Implementación de lógica en el backend para generar resúmenes ejecutivos basados en las métricas de los últimos ataques.', style='List Bullet')
    doc.add_paragraph('Sistema de Caché de Reportes: Mecanismo de caché (30 min) para evitar llamadas redundantes a la IA.', style='List Bullet')
    doc.add_paragraph('Threat Intel Loop: Proceso en segundo plano que analiza proactivamente IPs atacantes usando VirusTotal.', style='List Bullet')

    # Section 2
    doc.add_heading('2. Infraestructura de Tiempo Real (Webhooks & Wazuh)', level=1)
    doc.add_paragraph('Recepción de Webhooks: Creación del endpoint /api/webhook/wazuh para procesar alertas instantáneamente.', style='List Bullet')
    doc.add_paragraph('Sincronización Automática de Tickets: Alertas de severidad alta generan tickets automáticamente.', style='List Bullet')
    doc.add_paragraph('Configuración del Manager: Actualización de ossec.conf para habilitar integraciones externas.', style='List Bullet')

    # Section 3
    doc.add_heading('3. Sistema de Reportes Ejecutivos', level=1)
    doc.add_paragraph('Alineación de Datos: Sincronización de campos con los estándares del equipo (ej. Módulo de Julieta).', style='List Bullet')
    doc.add_paragraph('Métricas MITRE & Honeypot: Extracción automática de tácticas MITRE y contraseñas capturadas.', style='List Bullet')
    doc.add_paragraph('Opciones de Seguridad: Flexibilización de autenticación para acceso directivo.', style='List Bullet')

    # Section 4
    doc.add_heading('4. Optimizaciones de Frontend (Dashboard)', level=1)
    doc.add_paragraph('Reducción de Overhead: Optimización del polling en un 40%.', style='List Bullet')
    doc.add_paragraph('Sugerencia de Runbooks: Guías de respuesta integradas en alertas LSA.', style='List Bullet')
    doc.add_paragraph('Componentes HUD: Refactorización de AppCore.tsx para alto rendimiento.', style='List Bullet')

    # Section 5
    doc.add_heading('5. Arreglos y Estabilidad', level=1)
    doc.add_paragraph('Despliegue Flexible: Ejecución del frontend directamente fuera de Docker.', style='List Bullet')
    doc.add_paragraph('Sanitización: Limpieza de archivos temporales y preparación para producción.', style='List Bullet')
    doc.add_paragraph('Licencias: Incorporación de licencia GPLv2.', style='List Bullet')

    file_path = 'Integraciones_y_Arreglos_Valhalla_SOC.docx'
    doc.save(file_path)
    print(f'Documento creado exitosamente: {file_path}')

if __name__ == '__main__':
    create_doc()
