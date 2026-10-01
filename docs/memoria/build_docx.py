"""Genera la memoria en Word (.docx) editable a partir del mismo HTML que el PDF.

Títulos, párrafos, listas, tablas y capturas pasan a elementos nativos de Word; el
diagrama de arquitectura (Mermaid) se renderiza con el navegador y se inserta como imagen.

Uso:  python docs/memoria/build_docx.py [--grupo "Proyecto Valhalla"]
Requiere: pip install python-docx beautifulsoup4 playwright
"""
from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_memoria

HERE = Path(__file__).resolve().parent
OUT_DIR = None  # --out para generarlo en otra carpeta (p. ej. si el .docx está abierto en Word)
GREEN = RGBColor(0x0F, 0x5A, 0x3B)
ACCENT = RGBColor(0x1F, 0x8A, 0x5B)
MUTED = RGBColor(0x5D, 0x6E, 0x66)


def shade(cell_or_par, hex_fill: str) -> None:
    pr = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_runs(par, node, bold=False, italic=False, code=False, color=None, size=None) -> None:
    """Texto en línea con negritas, cursivas y código."""
    for ch in node.children if isinstance(node, Tag) else [node]:
        if isinstance(ch, NavigableString):
            text = re.sub(r"\s+", " ", str(ch))
            if not text:
                continue
            r = par.add_run(text)
            r.bold, r.italic = bold or None, italic or None
            if code:
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
            if color is not None:
                r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)
        elif isinstance(ch, Tag):
            if ch.name == "br":
                par.add_run().add_break()
            elif ch.name in ("b", "strong"):
                add_runs(par, ch, True, italic, code, color, size)
            elif ch.name in ("i", "em"):
                add_runs(par, ch, bold, True, code, color, size)
            elif ch.name == "code":
                add_runs(par, ch, bold, italic, True, color, size)
            elif ch.name == "span" and "muted" in (ch.get("class") or []):
                add_runs(par, ch, bold, italic, code, MUTED, size)
            else:
                add_runs(par, ch, bold, italic, code, color, size)


def add_table(doc, table: Tag) -> None:
    rows = table.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["th", "td"])) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Anchos proporcionales al texto (antes todas las columnas iguales: «ID» ocupaba un tercio)
    lens = [0] * ncols
    for tr in rows:
        for i, c in enumerate(tr.find_all(["th", "td"])):
            lens[i] = max(lens[i], min(len(c.get_text(" ", strip=True)), 120))
    total_cm, weights = 16.6, [max(l, 6) ** 0.8 for l in lens]
    widths = [max(1.8, total_cm * w / sum(weights)) for w in weights]
    t.autofit = False
    for tr in rows:
        cells = tr.find_all(["th", "td"])
        row_obj = t.add_row()
        trpr = row_obj._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")  # una fila no se parte entre páginas
        trpr.append(cant)
        row = row_obj.cells
        for i in range(ncols):
            row[i].width = Cm(widths[i])
        for i, c in enumerate(cells):
            p = row[i].paragraphs[0]
            is_head = c.name == "th"
            add_runs(p, c, bold=is_head, color=RGBColor(0xFF, 0xFF, 0xFF) if is_head else None)
            for r in p.runs:
                r.font.size = Pt(8.5) if not r.font.size else r.font.size
            if is_head:
                shade(row[i], "0F5A3B")
    doc.add_paragraph()


def add_image(doc, src: str, caption: str | None, width_cm: float) -> None:
    path = (HERE / src).resolve()
    if not path.exists():
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED


def mermaid_png(page_html: str) -> Path | None:
    """Renderiza el diagrama Mermaid con el navegador y devuelve un PNG."""
    from playwright.sync_api import sync_playwright
    tmp_html = HERE / "_docx_render.html"
    tmp_html.write_text(page_html, encoding="utf-8")
    out = Path(tempfile.gettempdir()) / "valhalla_arquitectura.png"
    try:
        with sync_playwright() as p:
            try:
                b = p.chromium.launch()
            except Exception:
                b = p.chromium.launch(channel="msedge")
            pg = b.new_page(viewport={"width": 1100, "height": 1400}, device_scale_factor=2)
            pg.goto(tmp_html.as_uri(), wait_until="networkidle")
            pg.wait_for_function("window.__mermaidDone === true", timeout=30000)
            el = pg.locator(".mermaid").first
            if el.count():
                el.screenshot(path=str(out))
            b.close()
    finally:
        tmp_html.unlink(missing_ok=True)
    return out if out.exists() else None


def build(grupo: str) -> Path:
    page = build_memoria.render(grupo)
    diagram = mermaid_png(page)
    soup = BeautifulSoup(page, "html.parser")

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        doc.styles[name].font.color.rgb = GREEN
        doc.styles[name].font.size = Pt(size)

    # ── Portada ──
    cover = soup.select_one(".cover")
    p = doc.add_paragraph()
    r = p.add_run(cover.select_one(".tag").get_text(strip=True))
    r.font.size, r.font.color.rgb, r.bold = Pt(9), ACCENT, True
    doc.add_paragraph().add_run().add_break()
    t = doc.add_paragraph()
    r = t.add_run(cover.h1.get_text(strip=True))
    r.font.size, r.bold, r.font.color.rgb = Pt(40), True, GREEN
    sub = doc.add_paragraph()
    r = sub.add_run(cover.select_one(".sub").get_text(strip=True))
    r.font.size, r.font.color.rgb = Pt(16), ACCENT
    desc = cover.find("p")
    if desc:
        dp = doc.add_paragraph()
        add_runs(dp, desc, size=11)
    for _ in range(6):
        doc.add_paragraph()
    meta = cover.select_one(".meta")
    mp = doc.add_paragraph()
    for ch in meta.children:
        if isinstance(ch, Tag) and "team" in (ch.get("class") or []):
            continue
        add_runs(mp, ch if isinstance(ch, Tag) else BeautifulSoup(str(ch), "html.parser"))
    for span in meta.select(".team span"):
        tp = doc.add_paragraph(style="List Bullet")
        add_runs(tp, span)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── Capítulos ──
    first = True
    for sec in soup.select("div.page > section"):
        if not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        for el in sec.children:
            if not isinstance(el, Tag):
                continue
            name = el.name
            cls = el.get("class") or []
            if name == "h2":
                doc.add_heading(el.get_text(strip=True), level=1)
            elif name == "h3":
                doc.add_heading(el.get_text(strip=True), level=2)
            elif name == "h4":
                doc.add_heading(el.get_text(strip=True), level=3)
            elif name == "p":
                add_runs(doc.add_paragraph(), el)
            elif name == "ul":
                for li in el.find_all("li", recursive=False):
                    add_runs(doc.add_paragraph(style="List Bullet"), li)
            elif name == "ol":
                # Numeración escrita en el texto: con el estilo «List Number» cada lista seguía la anterior
                for n, li in enumerate(el.find_all("li", recursive=False), 1):
                    lp = doc.add_paragraph()
                    lp.paragraph_format.left_indent = Cm(0.9)
                    lp.paragraph_format.first_line_indent = Cm(-0.6)
                    lp.add_run(f"{n}.  ")
                    add_runs(lp, li)
            elif name == "table":
                add_table(doc, el)
            elif name == "pre":
                pp = doc.add_paragraph()
                shade(pp, "EEF6F1")
                r = pp.add_run(el.get_text())
                r.font.name, r.font.size = "Consolas", Pt(8.5)
            elif name == "div" and "note" in cls:
                pp = doc.add_paragraph()
                shade(pp, "FDEEED" if "bad" in cls else "EEF6F1")
                add_runs(pp, el)
            elif name == "div" and "kpis" in cls:
                kt = doc.add_table(rows=1, cols=len(el.select(".kpi")))
                kt.style = "Table Grid"
                for i, k in enumerate(el.select(".kpi")):
                    c = kt.rows[0].cells[i].paragraphs[0]
                    rb = c.add_run(k.b.get_text(strip=True) + "\n")
                    rb.bold, rb.font.size, rb.font.color.rgb = True, Pt(16), GREEN
                    rs = c.add_run(k.span.get_text(strip=True))
                    rs.font.size, rs.font.color.rgb = Pt(8.5), MUTED
                doc.add_paragraph()
            elif name == "div" and "mermaid" in cls:
                if diagram:
                    doc.add_picture(str(diagram), width=Cm(16.5))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif name == "div" and "grid2" in cls:
                for fig in el.select("figure"):
                    img = fig.find("img")
                    cap = fig.find("figcaption")
                    phone = "phone" in (fig.get("class") or [])
                    add_image(doc, img["src"], cap.get_text(strip=True) if cap else None, 6.5 if phone else 15.5)
            elif name == "figure":
                img = el.find("img")
                cap = el.find("figcaption")
                add_image(doc, img["src"], cap.get_text(strip=True) if cap else None, 16)
            elif name == "ol" and "toc" in cls:
                pass

    # Pie de página con el título
    for s in doc.sections:
        fp = s.footer.paragraphs[0]
        fp.text = "Valhalla SOC · Memoria Práctica 3 · Grupo " + grupo
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = MUTED

    # Documento moderno (Word 2013+): sin esto Word lo abre en «Modo de compatibilidad»
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    for cs in compat.findall(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            compat.remove(cs)
    cs = OxmlElement("w:compatSetting")
    cs.set(qn("w:name"), "compatibilityMode")
    cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    cs.set(qn("w:val"), "15")
    compat.append(cs)

    out = Path(OUT_DIR or HERE) / f"P3_Grupo{grupo.replace(' ', '')}_Memoria.docx"
    doc.save(str(out))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--grupo", default="Proyecto Valhalla")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    OUT_DIR = a.out
    print(build(a.grupo))
